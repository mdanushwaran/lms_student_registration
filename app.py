import sqlite3
from pathlib import Path

import openpyxl
from flask import Flask, jsonify, request, send_file

app = Flask(__name__)

DB_FILE = Path("students.db")
EXCEL_FILE = Path("students.xlsx")


def init_db():
    conn = sqlite3.connect(DB_FILE)
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS students (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            reg_no TEXT NOT NULL,
            name TEXT NOT NULL,
            email TEXT NOT NULL,
            mobile TEXT NOT NULL,
            dept TEXT NOT NULL,
            year TEXT NOT NULL,
            batch TEXT NOT NULL,
            dob TEXT NOT NULL,
            gender TEXT NOT NULL,
            blood TEXT NOT NULL,
            address TEXT NOT NULL,
            username TEXT NOT NULL,
            photo TEXT,
            registered_at TEXT NOT NULL
        )
        """
    )
    conn.commit()
    conn.close()


def append_to_excel(student):
    headers = [
        "S.No",
        "Register No",
        "Student Name",
        "Email",
        "Mobile No",
        "Department",
        "Year",
        "Batch",
        "Date of Birth",
        "Gender",
        "Blood Group",
        "Address",
        "Username",
        "Registered At",
    ]

    if EXCEL_FILE.exists():
        workbook = openpyxl.load_workbook(EXCEL_FILE)
    else:
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)

    if "Student Registrations" not in workbook.sheetnames:
        sheet = workbook.create_sheet("Student Registrations")
    else:
        sheet = workbook["Student Registrations"]

    if sheet.max_row == 1 and all(cell.value is None for cell in sheet[1]):
        sheet.append(headers)

    row_number = sheet.max_row + 1
    sheet.append(
        [
            row_number - 1,
            student["regNo"],
            student["name"],
            student["email"],
            student["mobile"],
            student["dept"],
            student["year"],
            student["batch"],
            student["dob"],
            student["gender"],
            student["blood"],
            student["address"],
            student["username"],
            student["registeredAt"],
        ]
    )

    for cell in sheet[1]:
        cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        cell.fill = openpyxl.styles.PatternFill(fill_type="solid", fgColor="185FA5")
        cell.alignment = openpyxl.styles.Alignment(horizontal="center")

    sheet.column_dimensions["A"].width = 6
    sheet.column_dimensions["B"].width = 16
    sheet.column_dimensions["C"].width = 24
    sheet.column_dimensions["D"].width = 28
    sheet.column_dimensions["E"].width = 14
    sheet.column_dimensions["F"].width = 18
    sheet.column_dimensions["G"].width = 12
    sheet.column_dimensions["H"].width = 14
    sheet.column_dimensions["I"].width = 14
    sheet.column_dimensions["J"].width = 10
    sheet.column_dimensions["K"].width = 12
    sheet.column_dimensions["L"].width = 40
    sheet.column_dimensions["M"].width = 20
    sheet.column_dimensions["N"].width = 22

    workbook.save(EXCEL_FILE)


@app.before_request
def ensure_db():
    init_db()


@app.get("/")
def index():
    html = Path("lms_student_registration_form.html").read_text(encoding="utf-8")
    return html


@app.post("/register")
def register():
    data = request.get_json(silent=True) or {}

    student = {
        "regNo": (data.get("regNo") or "").strip().upper(),
        "name": (data.get("name") or "").strip(),
        "email": (data.get("email") or "").strip().lower(),
        "mobile": (data.get("mobile") or "").strip(),
        "dept": (data.get("dept") or "").strip(),
        "year": (data.get("year") or "").strip(),
        "batch": (data.get("batch") or "").strip(),
        "dob": (data.get("dob") or "").strip(),
        "gender": (data.get("gender") or "").strip(),
        "blood": (data.get("blood") or "").strip(),
        "address": (data.get("address") or "").strip(),
        "username": (data.get("username") or "").strip().lower(),
        "photo": (data.get("photo") or ""),
        "registeredAt": (data.get("registeredAt") or ""),
    }

    conn = sqlite3.connect(DB_FILE)
    conn.execute(
        """
        INSERT INTO students (
            reg_no, name, email, mobile, dept, year, batch, dob, gender,
            blood, address, username, photo, registered_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            student["regNo"],
            student["name"],
            student["email"],
            student["mobile"],
            student["dept"],
            student["year"],
            student["batch"],
            student["dob"],
            student["gender"],
            student["blood"],
            student["address"],
            student["username"],
            student["photo"],
            student["registeredAt"],
        ),
    )
    conn.commit()
    conn.close()

    append_to_excel(student)

    return jsonify({"ok": True, "download": "/download.xlsx"})


@app.get("/download.xlsx")
def download_excel():
    if not EXCEL_FILE.exists():
        return jsonify({"ok": False, "error": "No Excel file exists yet."}), 404
    return send_file(EXCEL_FILE, as_attachment=True, download_name="LMS_Students.xlsx")


if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=5000)
